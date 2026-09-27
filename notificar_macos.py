"""Genera avisos de evaluaciones atrasadas desde una planilla Excel en macOS.

Usa la aplicación Mail de macOS mediante AppleScript para crear mensajes en
borradores o enviarlos directamente desde la cuenta configurada. El modo
predeterminado solo muestra una vista previa en consola.
"""

from __future__ import annotations

import argparse
import re
import shlex
import subprocess
from datetime import date, datetime
from pathlib import Path

import openpyxl

CORREO_COPIA_FIJO = "crojas@colsanjavier.cl; guillermo.garcia@colsanjavier.cl; karen.ovalle@colsanjavier.cl"
CORREO_REMITENTE = "rolando.montero@colsanjavier.cl"
MESES = (
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
)
DIAS = ("lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo")
CORREO_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


def parsear_fecha(valor: str) -> date:
    """Acepta una fecha ISO (AAAA-MM-DD) o el formato DD/MM/AAAA."""
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(valor, formato).date()
        except ValueError:
            continue
    raise argparse.ArgumentTypeError("La fecha debe tener formato AAAA-MM-DD, DD/MM/AAAA o DD-MM-AAAA")


def limpiar_correo(valor: object) -> str | None:
    correo = str(valor or "").strip()
    return correo if CORREO_RE.fullmatch(correo) else None


def construir_mensaje(plantilla: str, fila: dict[str, object], fecha: date) -> str:
    tratamiento = "a" if str(fila.get("S", "")).strip().lower() == "a" else "o"
    reemplazos = {
        "[Alumno]": str(fila["Alumno"]).strip(),
        "[S]": tratamiento,
        "[fecha]": f"{DIAS[fecha.weekday()]} {fecha.day} de {MESES[fecha.month - 1]}",
        "[dd]": str(fecha.day),
        "[mes]": MESES[fecha.month - 1],
        "[Contenido]": str(fila["Contenido"]).strip() if fila.get("Contenido") else "   ",
    }
    for marcador, valor in reemplazos.items():
        plantilla = plantilla.replace(marcador, valor)
    return plantilla


def cargar_mensajes(ruta_excel: Path, ruta_plantilla: Path, fecha: date) -> list[dict[str, str]]:
    libro = openpyxl.load_workbook(ruta_excel, data_only=True, read_only=True)
    hoja = libro.active
    encabezados = [str(celda.value or "").strip() for celda in next(hoja.iter_rows())]
    requeridos = {"Alumno", "email alumno", "Mail Padre", "Mail Madre", "S", "PJ", "Curso", "Contenido"}
    faltantes = requeridos - set(encabezados)
    if faltantes:
        raise ValueError(f"Faltan columnas en la planilla: {', '.join(sorted(faltantes))}")

    plantilla = ruta_plantilla.read_text(encoding="utf-8")
    mensajes: list[dict[str, str]] = []

    for numero_fila, valores in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        fila = dict(zip(encabezados, valores))
        alumno = str(fila.get("Alumno") or "").strip()
        destinatario = limpiar_correo(fila.get("email alumno"))
        if not alumno or not destinatario:
            print(f"Fila {numero_fila}: omitida por alumno o correo principal inválido")
            continue

        copia: list[str] = []
        for columna in ("Mail Padre", "Mail Madre", "PJ"):
            correo = limpiar_correo(fila.get(columna))
            if correo and correo.lower() not in {item.lower() for item in copia}:
                copia.append(correo)
            elif fila.get(columna):
                print(f"Fila {numero_fila}: correo inválido omitido en {columna}: {fila[columna]}")

        if CORREO_COPIA_FIJO.lower() not in {item.lower() for item in copia}:
            copia.append(CORREO_COPIA_FIJO)

        mensajes.append({
            "alumno": alumno,
            "curso": str(fila.get("Curso") or "").strip(),
            "para": destinatario,
            "cc": "; ".join(copia),
            "cuerpo": construir_mensaje(plantilla, fila, fecha),
        })

    return mensajes


def _escapar_applescript(valor: str) -> str:
    return valor.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n").replace("\r", "")


def _script_correo(mensaje: dict[str, str], enviar: bool, remitente: str) -> str:
    asunto = f"Prueba atrasada | {mensaje['alumno']} | {mensaje['curso']}"
    destinatarios = [d.strip() for d in mensaje["para"].split(";") if d.strip()]
    ccs = [d.strip() for d in mensaje["cc"].split(";") if d.strip()]

    to_list = ", ".join(f'"{_escapar_applescript(d)}"' for d in destinatarios)
    cc_list = ", ".join(f'"{_escapar_applescript(d)}"' for d in ccs)
    asunto_esc = _escapar_applescript(asunto)
    cuerpo_esc = _escapar_applescript(mensaje["cuerpo"])

    script = f''' 
    tell application "Mail"
        set theMessage to make new outgoing message with properties {{subject:"{asunto_esc}", content:"{cuerpo_esc}"}}
        tell theMessage
            set visible to true
            set to recipients to {{{to_list}}}
            if {{{cc_list}}} is not {{}} then
                set cc recipients to {{{cc_list}}}
            end if
        end tell
        if "{_escapar_applescript(remitente)}" is not "" then
            try
                set targetAccount to first account whose email addresses contains "{_escapar_applescript(remitente)}"
                set sender of theMessage to targetAccount
            end try
        end if
        if {str(bool(enviar)).lower()} then
            send theMessage
        else
            save theMessage
        end if
    end tell
    '''
    return script.strip()


def enviar_con_mail(mensaje: dict[str, str], enviar: bool, borradores: bool, remitente: str) -> None:
    if not enviar and not borradores:
        print(f"\n--- {mensaje['alumno']} ---\n{mensaje['cuerpo']}\n")
        return

    script = _script_correo(mensaje, enviar=enviar, remitente=remitente)
    comando = ["osascript", "-e", script]
    resultado = subprocess.run(comando, capture_output=True, text=True)
    if resultado.returncode != 0:
        print(f"No se pudo {'enviar' if enviar else 'guardar'} el correo de {mensaje['alumno']}: {resultado.stderr.strip() or resultado.stdout.strip()}")
        return

    accion = "enviado" if enviar else "guardado como borrador"
    print(f"{accion}: {mensaje['alumno']}")


def ejecutar(mensajes: list[dict[str, str]], enviar: bool, borradores: bool, remitente: str) -> None:
    print(f"Preparando para {'enviar' if enviar else 'guardar borradores' if borradores else 'mostrar vista previa'}...\n")

    for mensaje in mensajes:
        if not enviar and not borradores:
            print(f"\n--- Prueba atrasada | {mensaje['alumno']} | {mensaje['curso']} ---\n")
            print(mensaje["cuerpo"])
            print("\n" + "=" * 40 + "\n")
            continue

        enviar_con_mail(mensaje, enviar=enviar, borradores=borradores, remitente=remitente)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fecha", type=parsear_fecha, required=True, help="Fecha de la evaluación, por ejemplo 2026-08-29")
    parser.add_argument("--excel", type=Path, default=Path("AUSENTE.xlsx"))
    parser.add_argument("--plantilla", type=Path, default=Path("carta_ausentes.txt"))
    parser.add_argument("--remitente", default=CORREO_REMITENTE, help=f"Cuenta de Mail que enviará los correos (predeterminada: {CORREO_REMITENTE})")
    grupo = parser.add_mutually_exclusive_group()
    grupo.add_argument("--enviar", action="store_true", help="Envía los correos usando Mail.app")
    grupo.add_argument("--borradores", action="store_true", help="Guarda los correos en Borradores de Mail")
    args = parser.parse_args()

    mensajes = cargar_mensajes(args.excel, args.plantilla, args.fecha)
    print(f"Mensajes preparados: {len(mensajes)}")
    ejecutar(mensajes, enviar=args.enviar, borradores=args.borradores, remitente=args.remitente)


if __name__ == "__main__":
    main()
