from pptx import Presentation
from servidor_mcp import aplicar_formato_markdown

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
body = slide.placeholders[1]
p = body.text_frame.add_paragraph()

texto = "Esto es **negrita**, esto es *cursiva* y esto es __negrita2__ y _cursiva2_ y texto normal."
aplicar_formato_markdown(p, texto)

print(f"Texto final: {p.text}")
for run in p.runs:
    print(f"Run: '{run.text}' | Bold: {run.font.bold} | Italic: {run.font.italic}")

# Test list case
tf = body.text_frame
tf.text = ""
points = ["**Punto 1** con *estilo*", "Normal"]
for pt in points:
    p_new = tf.add_paragraph()
    aplicar_formato_markdown(p_new, pt)

print("\nList test:")
for para in tf.paragraphs:
    print(f"Para: {para.text}")
    for run in para.runs:
        print(f"  Run: '{run.text}' | Bold: {run.font.bold} | Italic: {run.font.italic}")
