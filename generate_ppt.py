from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os

def create_presentation():
    # Setup
    output_folder = r"1ERO MEDIO FISICA"
    output_filename = "Clase20_OndasEstacionarias_1eroMedio.pptx"
    output_path = os.path.join(output_folder, output_filename)

    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    prs = Presentation()

    # Metadata
    metadata = {
        "NC": "20",
        "CURSO": "1ero Medio",
        "FECHA": "2026-07-01",
        "ASIGNATURA": "Física",
        "UNIDAD": "Ondas y Acústica",
        "CONTENIDOS": "Ondas estacionarias, nodos, antinodos, resonancia.",
        "HABILIDADES": "Analizar, comprender, identificar.",
        "VIRTUD": "Respeto"
    }

    # 13 Slide Content
    slides_content = [
        {
            "titulo": "Ondas Estacionarias",
            "cuerpo": [
                "Análisis de nodos, antinodos y resonancia",
                f"Curso: {metadata['CURSO']}",
                f"Clase: {metadata['NC']}",
                f"Asignatura: {metadata['ASIGNATURA']}"
            ]
        },
        {
            "titulo": "Objetivo de la Clase",
            "cuerpo": [
                "Analizar y comprender la formación y características de las ondas estacionarias, identificando nodos, antinodos, longitud de onda y condiciones de resonancia."
            ]
        },
        {
            "titulo": "Ruta de Clase",
            "cuerpo": [
                "1. Activación y Normas (Respeto)",
                "2. Introducción y Repaso",
                "3. Teoría de Ondas Estacionarias",
                "4. Nodos, Antinodos y λ",
                "5. Resonancia y Armónicos",
                "6. Práctica y Cierre"
            ]
        },
        {
            "titulo": "Activación Emocional y Normas",
            "cuerpo": [
                f"Virtud del Mes: {metadata['VIRTUD']}",
                "Norma: Escucha activa y respeto al turno de habla",
                "Reflexión: El respeto crea un ambiente seguro para preguntar y aprender."
            ]
        },
        {
            "titulo": "Inicio Motivacional",
            "cuerpo": [
                "¿Por qué una cuerda de guitarra vibra en patrones específicos y no de forma caótica?",
                "Descubriremos que existen ondas que parecen 'estacionarias' en el espacio."
            ]
        },
        {
            "titulo": "Repaso de Clase Anterior",
            "cuerpo": [
                "¿Qué ocurre cuando dos ondas se superponen?",
                "- Interferencia Constructiva: Mayor amplitud",
                "- Interferencia Destructiva: Menor o nula amplitud"
            ]
        },
        {
            "titulo": "Explicación y Modelamiento",
            "cuerpo": [
                "Formación: Superposición de dos ondas iguales en sentidos opuestos",
                "Nodos: Puntos de amplitud cero (estacionarios)",
                "Antinodos: Puntos de máxima oscilación",
                "Relación: λ/2 es la distancia entre dos nodos consecutivos"
            ]
        },
        {
            "titulo": "Práctica Guiada",
            "cuerpo": [
                "Analicemos un diagrama de cuerda vibrante",
                "Pregunta: Si una cuerda de 2m tiene 3 antinodos, ¿cuál es su longitud de onda?",
                "Paso 1: Identificar nodos",
                "Paso 2: Relacionar con λ"
            ]
        },
        {
            "titulo": "Práctica Independiente",
            "cuerpo": [
                "Resolución de problemas en el cuaderno:",
                "1. Dibujar el segundo armónico de una cuerda",
                "2. Calcular λ para una cuerda de 1.5m en el 1er armónico",
                "3. Identificar la condición de resonancia"
            ]
        },
        {
            "titulo": "Retroalimentación",
            "cuerpo": [
                "¿Confundimos nodos con antinodos?",
                "Recordar: Nodo = Quietud; Antinodo = Movimiento máximo",
                "Verificación de los cálculos de λ"
            ]
        },
        {
            "titulo": "Ticket de Salida",
            "cuerpo": [
                "Pregunta: En una onda estacionaria, ¿qué sucede con la amplitud en los nodos?",
                "Responde brevemente en tu hoja de salida."
            ]
        },
        {
            "titulo": "Reflexión y Cierre",
            "cuerpo": [
                "¿Cómo la sintonía de las ondas se parece al respeto mutuo?",
                "Conclusión: El equilibrio y la frecuencia adecuada generan armonía."
            ]
        },
        {
            "titulo": "Próxima Clase",
            "cuerpo": [
                "Ondas estacionarias en tubos sonoros y columnas de aire",
                "¡Preparen sus dudas sobre instrumentos de viento!"
            ]
        }
    ]

    for content in slides_content:
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)

        # Title
        title = slide.shapes.title
        title.text = content["titulo"]

        # Body
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.text = "" # Clear default

        for line in content["cuerpo"]:
            p = tf.add_paragraph()
            p.text = line
            p.space_after = Pt(10)

    # Add metadata to last slide or as footer (simplified: just a final slide)
    metadata_slide = prs.slides.add_slide(prs.slide_layouts[1])
    metadata_slide.shapes.title.text = "Información Institucional"
    tf_meta = metadata_slide.placeholders[1].text_frame
    tf_meta.text = ""
    for k, v in metadata.items():
        p = tf_meta.add_paragraph()
        p.text = f"{k}: {v}"

    prs.save(output_path)
    print(f"Successfully created {output_path}")

if __name__ == "__main__":
    create_presentation()
